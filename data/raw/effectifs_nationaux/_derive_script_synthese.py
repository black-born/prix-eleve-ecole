# -*- coding: utf-8 -*-
"""
Script de consolidation (fichier DERIVE, pas une source brute).
Lit les fichiers bruts officiels de ce dossier (sans les modifier) et produit
_derive_synthese_effectifs_nationaux_2023_2025.csv.

- Les valeurs lues dans les classeurs Excel (Notes d'Information DEPP) sont lues
  par programme (cellules indiquées dans la colonne "emplacement").
- Les valeurs issues de PDF (RERS, Etat de l'Ecole) sont saisies à la main ci-dessous
  avec leur emplacement exact (fiche / tableau / page) ; elles ont été vérifiées
  par extraction texte (pypdf) le 2026-10-04.
- Les lignes de type "calcul" sont des agrégats ou pondérations calculés ici :
  ce ne sont PAS des chiffres publiés.
Exécution : python _derive_script_synthese.py  (depuis ce dossier)
"""
import csv, os, unicodedata
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "_derive_synthese_effectifs_nationaux_2023_2025.csv")
rows = []

def add(indicateur, annee, valeur, champ, secteur, statut, fichier, emplacement, type_="publie", remarque=""):
    rows.append(dict(indicateur=indicateur, annee_rentree=annee, valeur=valeur, unite="élèves",
                     champ=champ, secteur=secteur, type=type_, statut=statut,
                     fichier_source=fichier, emplacement=emplacement, remarque=remarque))

def wb(name):
    return openpyxl.load_workbook(os.path.join(HERE, name), data_only=True)

def norm(s):
    s = str(s).strip().upper()
    return "".join(c for c in unicodedata.normalize("NFD", s) if unicodedata.category(c) != "Mn")

# ---------------------------------------------------------------- 1er degré
F58 = "depp_ni_2025-58_effectifs_1er_degre_rentree2025_donnees.xlsx"
ws = wb(F58)["Figure 7 en ligne"]
# colonnes : public F(2023) G(2024) H(2025) ; privé SC O P Q ; ensemble X Y Z
cols = {"public": {"2023": "F", "2024": "G", "2025": "H"},
        "prive_sous_contrat": {"2023": "O", "2024": "P", "2025": "Q"},
        "ensemble_public_prive_SC": {"2023": "X", "2024": "Y", "2025": "Z"}}
lignes = {"France (métropole + DROM, y.c. Mayotte)": 42, "France hors DROM (métropole)": 35,
          "DROM (y.c. Mayotte)": 41, "Mayotte": 40}
for lib, r in lignes.items():
    for sect, d in cols.items():
        for an, c in d.items():
            v = ws[f"{c}{r}"].value
            if v is None:
                continue
            add("1er degré - total", an, v, lib, sect, "définitif (constat de rentrée)", F58,
                f"onglet 'Figure 7 en ligne', cellule {c}{r}")

# détail 2024 et 2025 : NI 25.58 Figure 1 (C/D public, H/I privé SC, M/N ensemble)
ws = wb(F58)["Figure 1"]
# (la ligne TOTAL de la Figure 1, cellules C19..N19, est identique à 'Figure 7 en ligne' : non reprise ici)
det = {"préélémentaire": 10, "élémentaire (hors ULIS/UEEA)": 16, "UEEA": 17, "ULIS école": 18}
for lib, r in det.items():
    for sect, (c24, c25) in {"public": ("C", "D"), "prive_sous_contrat": ("H", "I"),
                             "ensemble_public_prive_SC": ("M", "N")}.items():
        for an, c in (("2024", c24), ("2025", c25)):
            add(f"1er degré - {lib}", an, ws[f"{c}{r}"].value, "France (métropole + DROM, y.c. Mayotte)",
                sect, "définitif (constat de rentrée)", F58, f"onglet 'Figure 1', cellule {c}{r}")
# détail 2023 : NI 23.50 Fig 1 (D public 2023-2024, I privé, N ensemble)
F50 = "depp_ni_2023-50_effectifs_1er_degre_rentree2023_donnees.xlsx"
ws = wb(F50)["Fig 1"]
# (ligne TOTAL D18/I18/N18 = 5 486 460 / 853 453 / 6 339 913, identique à NI 25.58 'Figure 7 en ligne' : non reprise)
for lib, r in {"préélémentaire": 10, "élémentaire (hors ULIS/UEEA)": 16, "ULIS école": 17}.items():
    for sect, c in {"public": "D", "prive_sous_contrat": "I", "ensemble_public_prive_SC": "N"}.items():
        add(f"1er degré - {lib}", "2023", ws[f"{c}{r}"].value, "France (métropole + DROM, y.c. Mayotte)",
            sect, "définitif (constat de rentrée)", F50, f"onglet 'Fig 1', cellule {c}{r}")

# hors contrat 1er degré : NI 25.58 Figure 6 en ligne (H21 = 2024, I21 = 2025) ; 2023 : RERS 2026 fiche 3.09
ws = wb(F58)["Figure 6 en ligne"]
add("1er degré - privé hors contrat (total)", "2024", ws["H21"].value, "France", "prive_hors_contrat",
    "définitif", F58, "onglet 'Figure 6 en ligne', cellule H21")
add("1er degré - privé hors contrat (total)", "2025", ws["I21"].value, "France", "prive_hors_contrat",
    "définitif", F58, "onglet 'Figure 6 en ligne', cellule I21")
add("1er degré - privé hors contrat (total)", "2023", 59477, "France", "prive_hors_contrat", "définitif",
    "rers2026_ch03_eleves_premier_degre.pdf", "fiche 3.09, tableau 'Évolution des effectifs du premier degré (1)', ligne Total, colonne 2023 (p. 91 de l'ouvrage)")
add("1er degré - élèves en classes hors contrat au sein d'écoles privées sous contrat (inclus dans privé SC)", "2025", 5469,
    "France", "prive_sous_contrat", "définitif", "rers2026_ch03_eleves_premier_degre.pdf",
    "fiche 3.09, note 1 du tableau 'Évolution des effectifs du premier degré (1)' (p. 91)")
add("1er degré - élèves en classes hors contrat au sein d'écoles privées sous contrat (inclus dans privé SC)", "2024", 6206,
    "France", "prive_sous_contrat", "définitif", "rers2025_ch03_eleves_premier_degre.pdf",
    "fiche 3.09, note 1 du tableau 'Évolution des effectifs du premier degré (1)' (p. 85)")

# ---------------------------------------------------------------- 2nd degré MEN
F59 = "depp_ni_2025-59_effectifs_2nd_degre_rentree2025_donnees.xlsx"
ws = wb(F59)["Figure 10 en ligne"]
# public C(2023) D(2024) E(2025) ; privé H I J ; ensemble M N O
cols2 = {"public": {"2023": "C", "2024": "D", "2025": "E"},
         "prive_sous_contrat": {"2023": "H", "2024": "I", "2025": "J"},
         "ensemble_public_prive_SC": {"2023": "M", "2024": "N", "2025": "O"}}
lig2 = {"1er cycle : formations en collège y.c. Segpa et ULIS": 13, "dont Segpa": 12, "dont ULIS collège": 10,
        "classe préparatoire à la seconde (créée en 2024)": 14,
        "2nd cycle professionnel (formations pro en lycée, y.c. ULIS pro)": 25, "dont ULIS pro": 24,
        "2nd cycle général et technologique (y.c. ULIS GT, hors prépa-seconde)": 30, "dont ULIS GT": 29,
        "formations en lycée y.c. ULIS et prépa-seconde": 31, "total": 32}
for lib, r in lig2.items():
    for sect, d in cols2.items():
        for an, c in d.items():
            v = ws[f"{c}{r}"].value
            if v is None:
                continue
            add(f"2nd degré MEN - {lib}", an, v, "France (métropole + DROM, y.c. Mayotte)", sect,
                "définitif (constat de rentrée, Sysca)", F59, f"onglet 'Figure 10 en ligne', cellule {c}{r}")

# DROM / métropole 2nd degré : somme des 5 académies DROM
DROM = {"GUADELOUPE", "MARTINIQUE", "GUYANE", "LA REUNION", "MAYOTTE"}
F42 = "depp_ni_2024-42_effectifs_2nd_degre_rentree2024_donnees.xlsx"
ws42 = wb(F42)["Compl2"]
ws11 = wb(F59)["Figure 11 en ligne"]
drom = {"2023": [0, 0, 0], "2024": [0, 0, 0], "2025": [0, 0, 0]}
may = {}
for r in ws42.iter_rows(values_only=True):
    if r[0] is not None and norm(r[0]) in DROM:
        drom["2023"][0] += r[1]; drom["2023"][1] += r[5]; drom["2023"][2] += r[9]
        if norm(r[0]) == "MAYOTTE":
            may["2023"] = (r[1], r[5], r[9])
for r in ws11.iter_rows(values_only=True):
    if r[0] is not None and norm(r[0]) in DROM:
        drom["2024"][0] += r[1]; drom["2024"][1] += r[5]; drom["2024"][2] += r[9]
        drom["2025"][0] += r[2]; drom["2025"][1] += r[6]; drom["2025"][2] += r[10]
        if norm(r[0]) == "MAYOTTE":
            may["2024"] = (r[1], r[5], r[9]); may["2025"] = (r[2], r[6], r[10])
tot2 = {}
ws10 = wb(F59)["Figure 10 en ligne"]
for an in ("2023", "2024", "2025"):
    tot2[an] = [ws10[f"{cols2['public'][an]}32"].value, ws10[f"{cols2['prive_sous_contrat'][an]}32"].value,
                ws10[f"{cols2['ensemble_public_prive_SC'][an]}32"].value]
    src = (F42 + " onglet 'Compl2'") if an == "2023" else (F59 + " onglet 'Figure 11 en ligne'")
    for i, sect in enumerate(["public", "prive_sous_contrat", "ensemble_public_prive_SC"]):
        add("2nd degré MEN - total", an, drom[an][i], "DROM (y.c. Mayotte)", sect, "définitif", src,
            "somme des lignes Guadeloupe, Martinique, Guyane, La Réunion, Mayotte", "calcul (somme)")
        add("2nd degré MEN - total", an, may[an][i], "Mayotte", sect, "définitif", src, "ligne Mayotte")
        add("2nd degré MEN - total", an, tot2[an][i] - drom[an][i], "France hors DROM (métropole)", sect,
            "définitif", F59 + " + " + src, "France ('Figure 10 en ligne' ligne 32) moins DROM", "calcul (différence)")

# ---------------------------------------------------------------- hors contrat 2nd degré (< 16 ans)
for an, v, f, loc in (("2023", 23174, "rers2026_ch04_eleves_second_degre.pdf", "fiche 4.24, tableau 2, ligne Ensemble, colonne 'Moins de 16 ans 2023' (p. 143)"),
                      ("2024", 24597, "rers2026_ch04_eleves_second_degre.pdf", "fiche 4.24, tableau 2, ligne Ensemble, colonne 'Moins de 16 ans 2024' (p. 143)"),
                      ("2025", 24758, "rers2026_ch04_eleves_second_degre.pdf", "fiche 4.24, tableau 2, ligne Ensemble, colonne 'Moins de 16 ans 2025' (p. 143)")):
    add("2nd degré - privé hors contrat (élèves de moins de 16 ans uniquement)", an, v, "France", "prive_hors_contrat",
        "définitif (estimation Sysca, champ < 16 ans)", f, loc)

# ---------------------------------------------------------------- agriculture (2nd degré, statut scolaire)
agri = {"2023": (49180, 87215, 136395, 322, "rers2024_ch04_eleves_second_degre.pdf", "RERS 2024 fiche 4.26, tableau 2, ligne 'Total second degré (1)' (p. 135)"),
        "2024": (49402, 89333, 138735, 318, "rers2025_ch04_eleves_second_degre.pdf", "RERS 2025 fiche 4.27, tableau 2, ligne 'Total second degré (1)' (p. 139)"),
        "2025": (49964, 90806, 140770, 318, "rers2026_ch04_eleves_second_degre.pdf", "RERS 2026 fiche 4.26, tableau 2, ligne 'Total second degré' (p. 147)")}
for an, (pu, pr, tt, dt, f, loc) in agri.items():
    add("2nd degré agricole (ministère de l'Agriculture) - total", an, pu, "France", "public", "définitif", f, loc)
    add("2nd degré agricole (ministère de l'Agriculture) - total", an, pr, "France", "prive_sous_et_hors_contrat", "définitif", f, loc)
    add("2nd degré agricole (ministère de l'Agriculture) - total", an, tt, "France", "ensemble", "définitif", f, loc,
        remarque=f"dont {dt} élèves en établissements sous double tutelle MEN/Agriculture (déjà comptés dans le 2nd degré MEN)")
    add("2nd degré agricole - sans doubles comptes avec le MEN", an, tt - dt, "France", "ensemble", "définitif", f,
        loc + " moins note 1 (double tutelle)", "calcul (différence)", "cohérent avec RERS 2026 fiche 1.02 (136,1 / 138,4 / 140,5 milliers)")

# ---------------------------------------------------------------- santé (établissements hospitaliers et médico-sociaux)
for an, tot, part in (("2023", 77802, 11186), ("2024", 78232, 11726), ("2025", 78632, 11921)):
    add("Élèves des établissements hospitaliers et médico-sociaux, hors scolarisation partagée (sans double compte)", an, tot - part,
        "France hors Mayotte", "ensemble", "définitif", "rers2026_ch01_systeme_educatif.pdf",
        "fiche 1.07, tableau 2 : ligne 'Scolarisés en étab. hospitaliers et médico-sociaux' moins 'dont scolarisation partagée' (p. 25)",
        "calcul (différence)", "cohérent avec RERS 2026 fiche 1.02 ligne 'Établissements spécialisés de la santé' (66,6 / 66,5 / 66,7 milliers)")

# ---------------------------------------------------------------- apprentis du secondaire (niveaux 3 et 4, tous ministères)
for an, n3, n4, f, loc, st in (("2023", 221060, 164568, "rers2026_ch06_apprentis.pdf", "fiche 6.01, tableau 2, lignes 'Total niveau 3' et 'Total niveau 4', colonne 2023 (p. 219)", "définitif (SIFA, au 31/12/2023)"),
                               ("2024", 221549, 170486, "rers2026_ch06_apprentis.pdf", "fiche 6.01, tableau 2, colonne 2024 (p. 219)", "définitif (SIFA, au 31/12/2024)"),
                               ("2025", 219666, 172271, "rers2026_ch06_apprentis.pdf", "fiche 6.09, tableau 'Effectifs d'apprentis en CFA par niveau', colonne 2025-2026 (p. 235)", "dernier chiffre publié (SIFA 2026, au 31/12/2025)")):
    add("Apprentis préparant un diplôme du 2nd degré (niveaux 3 et 4), tous ministères", an, n3 + n4, "France", "CFA publics et privés", st, f, loc,
        "calcul (somme niveau 3 + niveau 4)", "à compter À PART : pas des élèves sous statut scolaire")

# ---------------------------------------------------------------- totaux calculés
def val(ind, an, champ, sect):
    for r in rows:
        if r["indicateur"] == ind and r["annee_rentree"] == an and r["champ"] == champ and r["secteur"] == sect:
            return r["valeur"]
    raise KeyError((ind, an, champ, sect))
FR = "France (métropole + DROM, y.c. Mayotte)"
for an in ("2023", "2024", "2025"):
    p1 = val("1er degré - total", an, FR, "public"); s1 = val("1er degré - total", an, FR, "prive_sous_contrat")
    p2 = val("2nd degré MEN - total", an, FR, "public"); s2 = val("2nd degré MEN - total", an, FR, "prive_sous_contrat")
    men = p1 + s1 + p2 + s2
    ag = val("2nd degré agricole - sans doubles comptes avec le MEN", an, "France", "ensemble")
    hc = val("1er degré - privé hors contrat (total)", an, "France", "prive_hors_contrat") + \
         val("2nd degré - privé hors contrat (élèves de moins de 16 ans uniquement)", an, "France", "prive_hors_contrat")
    sa = [r["valeur"] for r in rows if r["indicateur"].startswith("Élèves des établissements hospitaliers") and r["annee_rentree"] == an][0]
    ap = [r["valeur"] for r in rows if r["indicateur"].startswith("Apprentis préparant") and r["annee_rentree"] == an][0]
    add("TOTAL A : 1er + 2nd degré MEN, public", an, p1 + p2, FR, "public", "définitif", "calcul", "somme", "calcul (somme)")
    add("TOTAL A : 1er + 2nd degré MEN, privé sous contrat", an, s1 + s2, FR, "prive_sous_contrat", "définitif", "calcul", "somme", "calcul (somme)")
    add("TOTAL A : 1er + 2nd degré MEN, public + privé sous contrat", an, men, FR, "ensemble", "définitif", "calcul", "somme", "calcul (somme)")
    add("TOTAL B : A + 2nd degré agricole (sans doubles comptes)", an, men + ag, "France", "ensemble", "définitif", "calcul", "somme", "calcul (somme)")
    add("TOTAL C : B + privé hors contrat (1er degré + 2nd degré < 16 ans)", an, men + ag + hc, "France", "ensemble", "définitif", "calcul", "somme", "calcul (somme)")
    add("TOTAL D : C + établissements hospitaliers et médico-sociaux (tous élèves sous statut scolaire, toutes tutelles)", an, men + ag + hc + sa, "France", "ensemble", "définitif", "calcul", "somme", "calcul (somme)")
    add("TOTAL E : D + apprentis du secondaire (= périmètre RERS 1.02 'Total élèves et apprentis')", an, men + ag + hc + sa + ap, "France", "ensemble", "définitif", "calcul", "somme", "calcul (somme)",
        "RERS 2026 fiche 1.02 : 12 667,8 / 12 579,8 / 12 459,5 milliers (écart <= 0,4 millier : arrondis et double tutelle agricole)")

# ---------------------------------------------------------------- pondération année civile (convention du Compte de l'éducation)
def civ(ind, champ, sect):
    for an_civ, (a1, a2) in {"2024": ("2023", "2024"), "2025": ("2024", "2025")}.items():
        v = round(2 / 3 * val(ind, a1, champ, sect) + 1 / 3 * val(ind, a2, champ, sect))
        add(f"ANNÉE CIVILE {an_civ} (2/3 rentrée {a1} + 1/3 rentrée {a2}) - {ind}", f"civile {an_civ}", v, champ, sect,
            "calcul", "depp_dossier206_2016_compte_education_methodes.pdf",
            "règle des 2/3-1/3 : Dossier DEPP n°206, chap. 2, encadré 'Les dépenses moyennes par élève ou étudiant', p. 35 (p. 37 du PDF)",
            "calcul (pondération)")
civ("1er degré - total", FR, "public"); civ("1er degré - total", FR, "prive_sous_contrat"); civ("1er degré - total", FR, "ensemble_public_prive_SC")
civ("2nd degré MEN - total", FR, "public"); civ("2nd degré MEN - total", FR, "prive_sous_contrat"); civ("2nd degré MEN - total", FR, "ensemble_public_prive_SC")
civ("TOTAL A : 1er + 2nd degré MEN, public", FR, "public")
civ("TOTAL A : 1er + 2nd degré MEN, privé sous contrat", FR, "prive_sous_contrat")
civ("TOTAL A : 1er + 2nd degré MEN, public + privé sous contrat", FR, "ensemble")
civ("TOTAL B : A + 2nd degré agricole (sans doubles comptes)", "France", "ensemble")
civ("2nd degré agricole (ministère de l'Agriculture) - total", "France", "ensemble")
civ("2nd degré agricole (ministère de l'Agriculture) - total", "France", "public")
civ("2nd degré agricole (ministère de l'Agriculture) - total", "France", "prive_sous_et_hors_contrat")

# ---------------------------------------------------------------- repères de périmètre (rentrée 2025, hors dénominateur principal)
CTX = "contexte (hors champ 'France')"
add("Outre-mer hors DROM (COM + Nouvelle-Calédonie) - 1er degré public + privé sous contrat", "2025", 32381 + 29537,
    "COM (Saint-Pierre-et-Miquelon, Polynésie française, Wallis-et-Futuna) + Nouvelle-Calédonie", "ensemble_public_prive_SC", "définitif",
    "rers2026_ch11_france_outre_mer.pdf", "fiche 11.01, tableau 'Effectifs du premier degré ... en outre-mer à la rentrée 2025' (suite), lignes COM (32 381) et Nouvelle-Calédonie (29 537) (p. 429)",
    "calcul (somme)", CTX)
add("Outre-mer hors DROM (COM + Nouvelle-Calédonie) - 2nd degré public + privé sous contrat", "2025", 29833 + 27075,
    "COM + Nouvelle-Calédonie", "ensemble_public_prive_SC", "définitif", "rers2026_ch11_france_outre_mer.pdf",
    "fiche 11.01, tableau 'Effectifs du second degré ... en outre-mer à la rentrée 2025' (suite), lignes COM (29 833) et Nouvelle-Calédonie (27 075) (p. 429)",
    "calcul (somme)", CTX)
add("Post-bac en lycée sous tutelle MEN : CPGE (statut scolaire)", "2025", 72055, "France", "public", "définitif",
    "rers2026_ch07_etudiants.pdf", "fiche 7.11, tableau 'Effectifs d'étudiants en CPGE par niveau ... 2025-2026', ligne 'Éducation nationale', colonne Public Total (p. 259)",
    remarque="enseignants payés sur la mission Enseignement scolaire ; comptés dans le supérieur par la DEPP")
add("Post-bac en lycée sous tutelle MEN : CPGE (statut scolaire)", "2025", 13318, "France", "prive_sous_et_hors_contrat", "définitif",
    "rers2026_ch07_etudiants.pdf", "fiche 7.11, même tableau, ligne 'Éducation nationale', colonne Privé Total (p. 259)")
add("Post-bac en lycée sous tutelle MEN-ESR : STS et assimilées (statut scolaire)", "2025", 151087, "France", "public", "définitif",
    "rers2026_ch07_etudiants.pdf", "fiche 7.12, tableau 'Effectifs d'étudiants en STS et assimilées selon la formation et le ministère de tutelle en 2025-2026', ligne 'Éducation nationale et Enseignement supérieur', colonne Public Total (p. 261)")
add("Post-bac en lycée sous tutelle MEN-ESR : STS et assimilées (statut scolaire)", "2025", 54263, "France", "prive_sous_et_hors_contrat", "définitif",
    "rers2026_ch07_etudiants.pdf", "fiche 7.12, même tableau, colonne Privé Total (p. 261)")
add("Post-bac en lycée agricole : BTSA (STS Agriculture, statut scolaire)", "2025", 14673, "France", "ensemble", "définitif",
    "rers2026_ch07_etudiants.pdf", "fiche 7.12, même tableau, ligne 'Agriculture' (public 9 894 + privé 4 779) (p. 261)")
add("Apprentis de niveaux 3 et 4 formés en convention avec un EPLE (au 31/12/2024)", "2024", 15222 + 18604, "France", "public (EPLE)", "définitif",
    "rers2026_ch06_apprentis.pdf", "fiche 6.08, tableau 'Effectifs d'apprentis en EPLE par niveau de formation en 2024-2025', lignes Niveau 3 (15 222) et Niveau 4 (18 604) (p. 233)",
    "calcul (somme)", "inclus dans les apprentis du secondaire ; utiles pour l'approche par établissement (lycées accueillant des apprentis)")
rows.append(dict(indicateur="EREA : nombre d'établissements (public / privé sous contrat)", annee_rentree="2025", valeur="76 / 1", unite="établissements",
                 champ="France", secteur="public / prive_sous_contrat", type="publie", statut="définitif",
                 fichier_source="rers2026_ch02_etablissements.pdf",
                 emplacement="fiche 2.04, tableau 'Évolution du nombre d'établissements du second degré', lignes EREA (p. 39)",
                 remarque="élèves d'EREA inclus dans le 2nd degré MEN mais non isolés dans les tableaux d'effectifs nationaux ; 926 classes publiques à 10,0 élèves/classe et 15 classes privées à 14,5 (p. 39)"))

with open(OUT, "w", newline="", encoding="utf-8-sig") as fh:
    w = csv.DictWriter(fh, fieldnames=list(rows[0].keys()), delimiter=";")
    w.writeheader(); w.writerows(rows)
print(len(rows), "lignes écrites dans", OUT)

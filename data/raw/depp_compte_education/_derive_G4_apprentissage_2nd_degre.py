# -*- coding: utf-8 -*-
"""
G4 - Part publique de l'apprentissage du 2nd degré (convention du Compte de l'éducation DEPP).

Fichier DÉRIVÉ (calculs), pas une source. Lit uniquement les fichiers bruts du dossier
data/raw/depp_compte_education/ (et effectifs_nationaux pour un contrôle) et écrit
_derive_G4_part_publique_apprentissage_2nd_degre.csv.

Année de référence : 2024 provisoire (édition DEPP de septembre 2025 = dernière édition qui isole
« l'apprentissage du second degré » ; les séries de septembre 2026 ne sont pas encore publiées).
Unités : millions d'euros courants (M€), sauf mention.
"""
import os
import csv
import openpyxl

D = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(D, "_derive_G4_part_publique_apprentissage_2nd_degre.csv")
rows = []


def put(id_, label, value, unit, source, note=""):
    rows.append({"id": id_, "libelle": label, "valeur": round(value, 3) if isinstance(value, float) else value,
                 "unite": unit, "source_ou_calcul": source, "note": note})
    return value


def wb(name):
    return openpyxl.load_workbook(os.path.join(D, name), data_only=True)


# ---------------------------------------------------------------- 1. Agrégats DEPP 2024p (NI 25.52)
ni = wb("depp_ni_2025-52_compte_education_2024_donnees.xlsx")
f5, f4 = ni["Figure 5"], ni["Figure 4"]
DIE = put("die_tot_2024p", "DIE totale 2024p", f5["B35"].value * 1000, "M€", "NI 25.52 xlsx, Figure 5, B35")
DIE2 = put("die_2nd_2024p", "DIE second degré (apprentissage compris) 2024p", f5["B32"].value * 1000, "M€",
           "NI 25.52 xlsx, Figure 5, B32")
sh = {k: f4[c].value / 100 for k, c in
      [("etat", "C32"), ("ct", "C33"), ("apu", "C34"), ("ent", "C35"), ("men", "C36")]}
PUB2 = put("pub_2nd_2024p", "Financement public initial du 2nd degré (État+CT+APU) 2024p",
           DIE2 * (sh["etat"] + sh["ct"] + sh["apu"]), "M€", "NI 25.52 xlsx, Figure 4, C32:C34 x B32 Figure 5")
E2 = put("ent_2nd_2024p", "Financement initial des entreprises (dont OPCO) au 2nd degré 2024p", DIE2 * sh["ent"], "M€",
         "NI 25.52 xlsx, Figure 4, C35 (4,51955 %) x DIE 2nd degré")

# ---------------------------------------------------------------- 2. Apprentissage du 2nd degré (séries chronologiques sept. 2025)
s = wb("depp_series_chrono_structure_die_par_niveau_366606.xlsx")["Structure DIE par niveau fin"]
part_app = None
for r in s.iter_rows():
    if r[1].value == "Apprentissage du second degré":
        part_app = [c.value for c in r if c.value is not None][-1] / 100
put("part_app_sec_2024p", "Part de l'apprentissage du 2nd degré dans la DIE 2024p (arrondie à 0,1 pt)", part_app * 100, "%",
    "Séries chrono DEPP 366606, onglet « Structure DIE par niveau fin », ligne « Apprentissage du second degré », col. 2024p")
APP_A = put("die_app_sec_2024p_partDIE", "DIE apprentissage 2nd degré = part x DIE", part_app * DIE, "M€", "calcul",
            "arrondi de la part : +/-0,05 pt, soit +/-99 M€")
dep = wb("depp_series_chrono_depense_par_eleve_par_niveau_481445.xlsx")["Dépense par élève, niveau"]
in_second = False
cout_app = None
for r in dep.iter_rows():
    v = r[1].value
    if v == "Second degré":
        in_second = True
    if v == "Enseignement supérieur":
        in_second = False
    if in_second and v == "dont Apprentissage*":
        cout_app = [c.value for c in r if c.value is not None][-1]
put("cout_apprenti_sec_2024p", "Dépense moyenne par apprenti du 2nd degré 2024p", cout_app, "€",
    "Séries chrono DEPP 481445, ligne « dont Apprentissage* » (second degré), col. 2024p")
# effectifs d'apprentis de niveaux 3 et 4 (RERS 2026, tableau 6.01, au 31 décembre)
st23 = 221060 + 164568
st24 = 221549 + 170486
EFF_APP = put("eff_app_sec_civil_2024", "Apprentis niveaux 3-4, année civile 2024 (2/3 x 31/12/2023 + 1/3 x 31/12/2024)",
              2 / 3 * st23 + 1 / 3 * st24, "apprentis",
              "RERS 2026 chap. 6, fiche 6.01 (stocks 31/12/2023 : 221 060 + 164 568 ; 31/12/2024 : 221 549 + 170 486)",
              "convention DEPP des effectifs en année civile")
APP_B = put("die_app_sec_2024p_coutxeff", "DIE apprentissage 2nd degré = dépense par apprenti x effectifs", cout_app * EFF_APP / 1e6,
            "M€", "calcul")

# ---------------------------------------------------------------- 3. Producteurs (RERS 2026, 10.04 tableau 2, financement final 2024p)
t = wb("depp_rers2026_10-04_producteurs_education_donnees.xlsx")["10.04 Tableau 2"]
def cfa(row):
    g = lambda c: t[f"{c}{row}"].value or 0.0
    return {"etat": g("D"), "ct": g("E"), "apu": g("F"), "men": g("G"), "ent": g("H"), "tot": g("I")}
CFA = [cfa(11), cfa(21), cfa(27)]
cfa_tot = sum(c["tot"] for c in CFA)
cfa_pub = sum(c["etat"] + c["ct"] + c["apu"] for c in CFA)
cfa_men = sum(c["men"] for c in CFA)
cfa_ent = sum(c["ent"] for c in CFA)
put("cfa_tot_2024p", "CFA (publics + privés subventionnés + non subventionnés), tous niveaux", cfa_tot, "M€",
    "RERS 2026 10.04 T2, I11+I21+I27")
put("cfa_pub_2024p", "dont État + CT + APU (financement final)", cfa_pub, "M€", "RERS 2026 10.04 T2, D+E+F lignes 11, 21, 27")
put("cfa_part_pub_2024p", "Part publique du financement des CFA, tous niveaux", cfa_pub / cfa_tot * 100, "%", "calcul")
put("cfa_ent_2024p", "dont entreprises et autres financeurs privés (OPCO)", cfa_ent, "M€", "RERS 2026 10.04 T2, H11+H21+H27")
put("cfa_men_2024p", "dont ménages", cfa_men, "M€", "RERS 2026 10.04 T2, G11+G21+G27")
NA_MAX = sum((t[f"H{r}"].value or 0.0) for r in (10, 15, 20, 26))
put("na_max_2024p", "Financement privé des producteurs du 2nd degré hors CFA (collèges-lycées publics et PSC, hors contrat, enseignement spécial)",
    NA_MAX, "M€", "RERS 2026 10.04 T2, H10+H15+H20+H26",
    "majorant du financement « entreprises » du 2nd degré hors apprentissage (inclut aussi STS/CPGE et GRETA)")

# ---------------------------------------------------------------- 4. Recoupement OPCO x NPEC par niveau
# Stocks d'apprentis au 31/12/2024 par niveau (RERS 2026, fiche 6.08, colonne « Ensemble des apprentis »)
stock = {"N3": 221549, "N4": 170486, "N5": 236930, "N6": 167738, "N7-8": 253262}
# NPEC moyens pondérés en vigueur (IGF-IGAS, revue de dépenses, mars 2024, annexe II, tableau 3)
npec = {"N3": 6630, "N4": 7699, "N5": 7932, "N6": 7751, "N7-8": 8428}
# Coût de revient 2024 par niveau (France compétences, RUF 2025, figure 5 p. 41)
cdr = {"N3": 7590, "N4": 8640, "N5": 8620, "N6": 8348, "N7-8": 9439}
sh_npec = sum(stock[k] * npec[k] for k in ("N3", "N4")) / sum(stock[k] * npec[k] for k in stock)
sh_cdr = sum(stock[k] * cdr[k] for k in ("N3", "N4")) / sum(stock[k] * cdr[k] for k in stock)
put("part_sec_npec", "Part des niveaux 3-4 dans les NPEC (stocks 31/12/2024 x NPEC moyens pondérés)", sh_npec * 100, "%", "calcul")
put("part_sec_cout", "Part des niveaux 3-4 dans les coûts des OFA (stocks x coût de revient 2024)", sh_cdr * 100, "%", "calcul")
OPCO = 8710.0  # Jaune « Formation professionnelle » PLF 2026, tableau 3 : Opco, apprentis, 8,71 Md€ (2024p)
E_C1 = put("ent_app_sec_C1", "Financement OPCO de l'apprentissage niveaux 3-4 (variante C1 : OPCO DARES x part NPEC)", OPCO * sh_npec,
           "M€", "Jaune PLF 2026 tableau 3 (8,71 Md€) x part_sec_npec")
E_C2 = put("ent_app_sec_C2", "Variante C2 : entreprises des CFA (DEPP) x part des coûts de revient", cfa_ent * sh_cdr, "M€",
           "RERS 10.04 T2 x part_sec_cout")
M_SEC = put("men_app_sec", "Ménages, apprentissage 2nd degré (CFA ménages x part des coûts niveaux 3-4)", cfa_men * sh_cdr, "M€", "calcul")

# ---------------------------------------------------------------- 5. Part publique P de l'apprentissage du 2nd degré
for tag, APP in (("cxe", APP_B), ("part", APP_A)):
    put(f"P_min_{tag}", f"Borne basse : tout le financement « entreprises » du 2nd degré va à l'apprentissage (DIE app = {APP:.0f})",
        APP - E2 - cfa_men, "M€", "DIE_app - E_2nd - ménages CFA (tous niveaux)")
    put(f"P_max_{tag}", f"Borne haute : tout le privé hors CFA du 2nd degré est hors apprentissage (DIE app = {APP:.0f})",
        APP - (E2 - NA_MAX), "M€", "DIE_app - (E_2nd - na_max), ménages = 0")
    put(f"P_C1_{tag}", f"Estimation centrale C1 (DIE app = {APP:.0f})", APP - E_C1 - M_SEC, "M€", "DIE_app - ent_app_sec_C1 - men_app_sec")
    put(f"P_C2_{tag}", f"Estimation centrale C2 (DIE app = {APP:.0f})", APP - E_C2 - M_SEC, "M€", "DIE_app - ent_app_sec_C2 - men_app_sec")
put("NA_implicite_C1", "Financement « entreprises » du 2nd degré hors apprentissage impliqué par C1", E2 - E_C1, "M€",
    "à comparer au solde de la taxe d'apprentissage : 522 M€ en 2025, tous niveaux (SOLTéA)")
put("P_naif_prorata", "Prorata naïf : part publique des CFA tous niveaux x DIE app (incohérent avec E_2nd)", cfa_pub / cfa_tot * APP_A,
    "M€", "calcul", "implique un financement entreprises de l'apprentissage > E_2nd : à écarter comme estimation")

# ---------------------------------------------------------------- 6. Effet sur la dépense publique par élève sous statut scolaire
EFF2 = put("eff_2nd_implicite_2024p", "Effectif implicite du 2nd degré 2024p (DIE / 11 660 €)", DIE2 * 1e6 / 11660, "élèves",
           "NI 25.52 xlsx Figure 5 B32 / dépense moyenne 11 660 € (RERS 2026 10.05 T2)")
EFF_SCO = EFF2 - EFF_APP
put("pub_par_eleve_2nd_avec_apprentis", "Dépense publique par élève du 2nd degré, apprentis compris", PUB2 * 1e6 / EFF2, "€", "calcul")
for P in (300.0, 500.0, 650.0, 850.0, 1300.0):
    put(f"pub_par_eleve_scolaire_P{int(P)}", f"Dépense publique par élève sous statut scolaire si P = {P/1000:.2f} Md€",
        (PUB2 - P) * 1e6 / EFF_SCO, "€", "(pub_2nd - P) / (eff_2nd - eff_app)")

# Parts publiques s = P / DIE_app correspondantes
for tag, P, APP in (("min", APP_B - E2 - cfa_men, APP_B), ("C2_cxe", APP_B - E_C2 - M_SEC, APP_B),
                    ("C1_part", APP_A - E_C1 - M_SEC, APP_A), ("max", APP_A - (E2 - NA_MAX), APP_A)):
    put(f"s_{tag}", f"Part publique s de l'apprentissage du 2nd degré ({tag})", P / APP * 100, "%", "P / DIE_app")
S_CENTRAL = 0.175
put("s_central_retenu", "Part publique centrale retenue (milieu des variantes C1/C2)", S_CENTRAL * 100, "%",
    "hypothèse H-G4-2 : 16 à 19 % selon les variantes, bornes 8 à 35 %")
# Variante dénominateur « apprentis implicites » (DIE_app / 9 500 €), comme au § 11.6
EFF_APP_IMPL = APP_A * 1e6 / cout_app
put("pub_par_eleve_scolaire_s175_implicite", "Dépense publique par élève sous statut scolaire 2024p, s = 17,5 %, apprentis implicites",
    (PUB2 - S_CENTRAL * APP_A) * 1e6 / (EFF2 - EFF_APP_IMPL), "€", "(pub_2nd - s x 3 745) / (eff_2nd - 394 199)")
put("pub_par_eleve_scolaire_s175_sifa", "Idem, apprentis SIFA pondérés (387 764) et DIE_app = 3 684 M€",
    (PUB2 - S_CENTRAL * APP_B) * 1e6 / EFF_SCO, "€", "(pub_2nd - s x 3 684) / (eff_2nd - 387 764)")

# ---------------------------------------------------------------- 7. Transposition 2025p (NI 26.42), hypothèses du § 11.7 (H-a)
ni26 = wb("depp_ni_2026-42_compte_education_2025_donnees.xlsx")
DIE2_25 = ni26["Figure 5"]["B33"].value * 1000
f4b = ni26["Figure 4"]
PUB2_25 = DIE2_25 * sum(f4b[c].value for c in ("C32", "C33", "C34")) / 100
E2_25 = put("ent_2nd_2025p", "Financement initial des entreprises (dont OPCO) au 2nd degré 2025p", DIE2_25 * f4b["C35"].value / 100,
            "M€", "NI 26.42 xlsx, Figure 4, C35 x Figure 5, B33")
APP_25 = DIE2_25 * (APP_A / DIE2)  # H-a : même part de l'apprentissage dans la DIE du 2nd degré qu'en 2024p
EFF2_25 = DIE2_25 * 1e6 / 11780
EFF_APP_25 = 2 / 3 * 392035 + 1 / 3 * 391937  # SIFA 31/12/2024 et 31/12/2025 (RERS 2026, 6.09)
put("die_app_sec_2025p_Ha", "DIE apprentissage 2nd degré 2025p (H-a : part de 2024p dans le 2nd degré)", APP_25, "M€", "calcul")
put("P_min_2025p", "Borne basse 2025p (même raisonnement qu'en 2024p : ménages = total CFA 2024p, 82,9 M€)",
    APP_25 - E2_25 - cfa_men, "M€", "DIE_app_2025 - E_2nd_2025 - ménages CFA",
    "plus basse qu'en 2024p car le financement entreprises du 2nd degré augmente de 117 M€")
for s_ in (0.08, 0.175, 0.35):
    put(f"pub_par_eleve_scolaire_2025p_s{int(s_*1000)}", f"Dépense publique initiale par élève sous statut scolaire 2025p, s = {s_*100:.1f} %",
        (PUB2_25 - s_ * APP_25) * 1e6 / (EFF2_25 - EFF_APP_25), "€", "(pub_2nd_2025 - s x DIE_app_2025) / (N_2025 - 392 002)")

with open(OUT, "w", newline="", encoding="utf-8-sig") as fh:
    w = csv.DictWriter(fh, fieldnames=["id", "libelle", "valeur", "unite", "source_ou_calcul", "note"], delimiter=";")
    w.writeheader()
    w.writerows(rows)
for r in rows:
    print(f'{r["id"]:<38} {r["valeur"]!s:>14} {r["unite"]}')

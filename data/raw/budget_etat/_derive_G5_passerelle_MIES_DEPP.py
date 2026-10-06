# -*- coding: utf-8 -*-
"""
Tache G5 - Passerelle entre la mission « Enseignement scolaire » (MIES, budget de l'Etat)
et la part « Etat » du compte de l'education de la DEPP (financement initial, 1er + 2nd degres).

Script reproductible : lit uniquement des fichiers locaux (dossier budget_etat et ../depp_compte_education).
Sortie : _derive_G5_passerelle_MIES_DEPP_2024_2025.csv (meme dossier).
Tous les montants en millions d'euros courants (MEUR), CP executes, CAS Pensions compris.

Conventions (detail dans NOTES.md, section 14) :
- DEPP « Etat » = MEN-MESR + autres ministeres + reste du monde ; financement initial (bourses comprises).
- Champ DEPP = France (metropole + DROM) : les COM sont retirees.
- Classement DEPP documente : post-bac en lycee -> superieur ; formation continue des adultes et formation continue
  des personnels d'education -> extrascolaire ; allocations PFMP -> exclues de la DIE ; credits transversaux
  (remplacement, etc.) -> repartis entre niveaux par des cles (Dossier 206, p.32 imprimee).
"""
import os
import csv
import warnings
import pandas as pd
import openpyxl

warnings.filterwarnings("ignore")
HERE = os.path.dirname(os.path.abspath(__file__))
EXT = os.path.join(HERE, "extractions")
DEPP = os.path.normpath(os.path.join(HERE, "..", "depp_compte_education"))
M = 1e6

# ---------------------------------------------------------------- lectures
act = pd.read_csv(os.path.join(EXT, "rap2025_execution_par_action_2024_2025.csv"), sep=";", encoding="utf-8-sig")
prog = pd.read_csv(os.path.join(EXT, "rap_execution_2024_2025_par_programme_et_titre.csv"), sep=";", encoding="utf-8-sig")
inp = pd.read_csv(os.path.join(EXT, "passerelle_G5_intrants_RAP_DPT_2024_2025.csv"), sep=";", encoding="utf-8-sig")
tab2i = pd.read_csv(os.path.join(EXT, "depp_dossier206_tableau2i_financement_initial_etat_par_niveau_2013_2014.csv"), sep=";", encoding="utf-8-sig")
plr13 = pd.read_csv(os.path.join(EXT, "plr2013_MIES_execution_CP_par_programme_action.csv"), sep=";", encoding="utf-8-sig", dtype={"Programme": str, "Action": str})
dger = pd.read_csv(os.path.join(EXT, "dger_effectifs_enseignement_agricole_voie_scolaire_par_niveau_2023_2025.csv"), sep=";", encoding="utf-8-sig")

I = dict(zip(inp["id"], inp["valeur"]))


def A(p, a, col):  # montant d'une action (EUR -> MEUR)
    r = act[(act.programme == p) & (act.action == a)]
    assert len(r) == 1, (p, a)
    return float(r[col].iloc[0]) / M


def P(p, year, col="CP_consommes"):
    r = prog[(prog.programme.astype(str) == str(p)) & (prog.annee == year)]
    assert len(r) == 1, (p, year)
    return float(r[col].iloc[0]) / M


def cell(path, sheet, ref):
    wb = openpyxl.load_workbook(path, data_only=True)
    return wb[sheet][ref].value


ni2642 = os.path.join(DEPP, "depp_ni_2026-42_compte_education_2025_donnees.xlsx")
ni2552 = os.path.join(DEPP, "depp_ni_2025-52_compte_education_2024_donnees.xlsx")
chrono = os.path.join(DEPP, "depp_series_chrono_die_financeur_initial_final_par_niveau_366603.xlsx")
rers1003 = os.path.join(DEPP, "depp_rers2026_10-03_budget_donnees.xlsx")

# DEPP 2025p (NI 26.42)
etat_part_1_25 = cell(ni2642, "Figure 4", "B32") / 100
etat_part_2_25 = cell(ni2642, "Figure 4", "C32") / 100
die1_25 = cell(ni2642, "Figure 5", "B32") * 1000
die2_25 = cell(ni2642, "Figure 5", "B33") * 1000
menesr_tous_niveaux_25 = cell(ni2642, "Figure 3bis", "F5")
menesr_tous_niveaux_24def = cell(ni2642, "Figure 3bis", "E5")
# DEPP 2024p (NI 25.52 + series chronologiques de septembre 2025)
etat_part_1_24 = cell(ni2552, "Figure 4", "B32") / 100
etat_part_2_24 = cell(ni2552, "Figure 4", "C32") / 100
die1_24 = cell(ni2552, "Figure 5", "B31") * 1000
die2_24 = cell(ni2552, "Figure 5", "B32") * 1000
men_part_1_24 = cell(chrono, "DIE 1er degré", "U17") / 100   # « dont MENESR », arrondi a 0,1 point
men_part_2_24 = cell(chrono, "DIE 2nd degré", "U17") / 100
# MIRES (RERS 2026, fiche 10.03, tableau 3 ; champ France + COM ; source RAP)
wb = openpyxl.load_workbook(rers1003, data_only=True)
ws = wb["10.03 Tableau 3"]
mires = {}  # valeurs deja en MEUR dans la RERS (cellules A18 et A26, colonnes C = 2024, D = 2025)
for r in range(1, ws.max_row + 1):
    lab = ws.cell(r, 1).value
    lab = lab.replace("\xa0", " ") if isinstance(lab, str) else lab
    if isinstance(lab, str) and lab.startswith("Programme 150"):
        mires["P150"] = {2024: ws.cell(r, 3).value, 2025: ws.cell(r, 4).value}
    if isinstance(lab, str) and lab.startswith("Programme 231"):
        mires["P231"] = {2024: ws.cell(r, 3).value, 2025: ws.cell(r, 4).value}
assert "P150" in mires and "P231" in mires

# ---------------------------------------------------------------- budget MIES
MEN_PROGS = [140, 141, 230, 139, 214]
mies = {y: P("MISSION", y) for y in (2024, 2025)}
men = {y: sum(P(p, y) for p in MEN_PROGS) for y in (2024, 2025)}
p143 = {y: P(143, y) for y in (2024, 2025)}
com_men_24 = sum(I[f"com_p{p}_2024"] for p in MEN_PROGS) / M
com_143_24 = I["com_p143_2024"] / M
com_mies_24 = com_men_24 + com_143_24
com_mires_24 = (I["com_p150_pf_2024"] + I["com_p150_nc_2024"] + I["com_p231_pf_2024"] + I["com_p231_nc_2024"]) / M

postbac = {2024: A(141, 5, "CP_conso_2024") + A(139, 6, "CP_conso_2024"),
           2025: A(141, 5, "CP_conso_2025") + A(139, 6, "CP_conso_2025")}
fc_adultes = {2024: A(141, 9, "CP_conso_2024"), 2025: A(141, 9, "CP_conso_2025")}
apprentissage = {2024: A(141, 4, "CP_conso_2024"), 2025: A(141, 4, "CP_conso_2025")}
p214_11 = {2024: A(214, 11, "CP_conso_2024"), 2025: A(214, 11, "CP_conso_2025")}
pfmp_men = {2024: (I["pfmp_p141_2024"] + I["pfmp_p139_2024"]) / M, 2025: (I["pfmp_p141_2025"] + I["pfmp_p139_2025"]) / M}
pfmp_143 = {2024: I["pfmp_p143_2024"] / M, 2025: I["pfmp_p143_2025"] / M}
meef = {2025: (I["meef_p140_2025"] + I["meef_p141_2025"] + I["meef_p139_2025"] + I["meef_p230_2025"]) / M,
        # 2024 : P140 non detaille dans le RAP 2024 -> valeur 2025 du P140 reprise (hypothese)
        2024: (I["meef_p140_2025"] + I["meef_p141_2024"] + I["meef_p139_2024"] + I["meef_p230_2024"]) / M}
formation_actions = {y: A(140, 4, f"CP_conso_{y}") + A(141, 10, f"CP_conso_{y}") + A(139, 10, f"CP_conso_{y}") for y in (2024, 2025)}
formation_HT2_25 = A(140, 4, "CP_conso_2025_hors_titre2") + A(141, 10, "CP_conso_2025_hors_titre2") + A(139, 10, "CP_conso_2025_hors_titre2")
formation_T2_25 = A(140, 4, "CP_conso_2025_titre2") + A(141, 10, "CP_conso_2025_titre2") + A(139, 10, "CP_conso_2025_titre2")

# stagiaires (categorie 1108) : ETPT x cout moyen HCAS ; estimation avec CAS par le ratio CAS / T2 hors CAS du programme
def ratio_cas(p, y=2025):
    return P(p, y, "CP_dont_CAS_pensions") / P(p, y, "CP_titre2_hors_CAS")
stag = {}
for p in (140, 141, 230, 139):
    etpt = I[f"stag_etpt_p{p}_2025"]
    cout = I[f"stag_cout_p{p}_2025"]
    hcas = etpt * cout / M
    r = ratio_cas(141) if p == 230 else ratio_cas(p)  # P230 : ratio du P141 (stagiaires CPE fonctionnaires)
    stag[p] = (etpt, cout, hcas, hcas * (1 + r))

# ---------------------------------------------------------------- effectifs agricoles (annee civile 2/3 - 1/3)
def eff(an, niv):
    return float(dger[(dger.annee_scolaire == an) & (dger.niveau == niv)].effectif_eleves_etudiants.iloc[0])
AN = {2023: "2023-2024 (rentrée 2023)", 2024: "2024-2025 (rentrée 2024)", 2025: "2025-2026 (rentrée 2025)"}
sup = {k: eff(v, "BTS") + eff(v, "BTSA") + eff(v, "CPGE") for k, v in AN.items()}
tot = {k: eff(v, "Total") for k, v in AN.items()}
part_btsa = {y: (2 / 3 * sup[y - 1] + 1 / 3 * sup[y]) / (2 / 3 * tot[y - 1] + 1 / 3 * tot[y]) for y in (2024, 2025)}
# variante ponderee par secteur (cout par eleve public / prive approche par les actions 01 et 02 du P143, 2025)
# effectifs 2025 public / prive par niveau : DGER onglet 3, voir NOTES ; valeurs reprises ici
pub_sup25, pub_tot25, priv_sup25, priv_tot25 = 10566, 61118, 4779, 96268
cout_pub = A(143, 1, "CP_conso_2025") / pub_tot25
cout_priv = A(143, 2, "CP_conso_2025") / priv_tot25
w_sup = pub_sup25 * cout_pub + priv_sup25 * cout_priv
w_tot = pub_tot25 * cout_pub + priv_tot25 * cout_priv
part_btsa_pond_25 = w_sup / w_tot
# variante structure DEPP 2013 (Agriculture -> superieur court / (2nd degre + superieur court))
t13 = tab2i[(tab2i.annee == 2013)].set_index("financeur_initial")
agri13 = t13.loc["État - Agriculture"]
part_btsa_depp13 = agri13["sup_technique_court_MEUR"] / (agri13["total_2nd_degre_MEUR"] + agri13["sup_technique_court_MEUR"])

# ---------------------------------------------------------------- calage 2013 (Dossier 206)
es13 = t13.loc["État - Enseignement scolaire"]
plr13["CP_total_EUR"] = plr13["CP_total_EUR"].astype(float)
def plr(p, a=None):
    x = plr13[plr13.Programme == str(p)]
    if a is not None:
        x = x[x.Action == f"{a:02d}"]
    return x.CP_total_EUR.sum() / M
men13_budget = sum(plr(p) for p in MEN_PROGS)
postbac13 = plr(141, 5) + plr(139, 6)
fc13 = plr(141, 9)
hors12_13 = es13["total_superieur_MEUR"] + es13["total_extrascolaire_MEUR"]
non_ident_13 = hors12_13 - postbac13 - fc13
part_non_ident_13 = non_ident_13 / es13["total_MEUR"]
part_hors12_13 = hors12_13 / es13["total_MEUR"]
autres_rdm_12_13 = (t13.loc["État - Autres ministères"]["total_1er_degre_MEUR"] + t13.loc["État - Autres ministères"]["total_2nd_degre_MEUR"]
                    + t13.loc["Reste du monde"]["total_1er_degre_MEUR"] + t13.loc["Reste du monde"]["total_2nd_degre_MEUR"])
agri13_2nd_share = agri13["total_2nd_degre_MEUR"] / plr(143)

# ---------------------------------------------------------------- 2024 : decomposition observee (DEPP 2024p)
etat12_24 = etat_part_1_24 * die1_24 + etat_part_2_24 * die2_24
men12_24 = men_part_1_24 * die1_24 + men_part_2_24 * die2_24
men_france_24 = men[2024] - com_men_24
hors12_men_24 = men_france_24 - men12_24
ident_men_24 = postbac[2024] + fc_adultes[2024] + pfmp_men[2024] + p214_11[2024] + meef[2024]
non_ident_24 = hors12_men_24 - ident_men_24
part_non_ident_24 = non_ident_24 / men_france_24
p143_france_24 = p143[2024] - com_143_24
btsa_24 = part_btsa[2024] * (p143_france_24 - pfmp_143[2024])
p143_2nd_24 = p143_france_24 - pfmp_143[2024] - btsa_24
autres_rdm_12_24 = etat12_24 - men12_24 - p143_2nd_24

# controle global MEN-MESR (tous niveaux)
glob = {}
for y in (2024, 2025):
    b = men[y] + mires["P150"][y] + mires["P231"][y] - com_men_24 - com_mires_24
    glob[y] = (b, b - pfmp_men[y])
depp_glob = {2024: menesr_tous_niveaux_24def, 2025: menesr_tous_niveaux_25}

# ---------------------------------------------------------------- 2025 : passerelle
etat12_25 = etat_part_1_25 * die1_25 + etat_part_2_25 * die2_25
men_france_25 = men[2025] - com_men_24
p143_france_25 = p143[2025] - com_143_24
btsa_25 = part_btsa[2025] * (p143_france_25 - pfmp_143[2025])
france_25 = mies[2025] - com_mies_24
ident_25 = postbac[2025] + fc_adultes[2025] + pfmp_men[2025] + pfmp_143[2025] + p214_11[2025] + meef[2025] + btsa_25
etape1 = france_25 - ident_25
c1_central = part_non_ident_24 * men_france_25
c1_var13 = part_non_ident_13 * men_france_25
etape2 = etape1 - c1_central
etape3 = etape2 + autres_rdm_12_24
residu_central = etape3 - etat12_25
# variante : calage 2013 complet (structure du Dossier 206)
var13 = (men_france_25 * (1 - part_hors12_13) - (pfmp_men[2025] + p214_11[2025] + meef[2025])
         + (p143[2025] - pfmp_143[2025]) * agri13_2nd_share + autres_rdm_12_13)
# variante : calage 2024 mais BTSA au prorata pondere ou structure DEPP 2013
def bridge_with_btsa(share):
    b25 = share * (p143_france_25 - pfmp_143[2025])
    b24 = share * (p143_france_24 - pfmp_143[2024])
    autres = etat12_24 - men12_24 - (p143_france_24 - pfmp_143[2024] - b24)
    return france_25 - (ident_25 - btsa_25 + b25) - c1_central + autres
var_btsa_pond = bridge_with_btsa(part_btsa_pond_25) - etat12_25
var_btsa_d13 = bridge_with_btsa(part_btsa_depp13) - etat12_25
# variante : COM 2025 = LFI 2025 (1 103,5 MEUR) au lieu de l'execution 2024
com_lfi25 = 1103505942 / M
var_com = residu_central + (com_mies_24 - com_lfi25) * (1 - 0)  # la part MEN des COM entre aussi dans C1 (effet de second ordre neglige)
# ancienne passerelle (NOTES § 8)
ancienne = mies[2025] - com_mies_24 - postbac[2025] - p214_11[2025] - fc_adultes[2025] - apprentissage[2025]

# ---------------------------------------------------------------- sortie
rows = []
def add(annee, bloc, poste, montant, classement, statut, source, calcul=""):
    rows.append([annee, bloc, poste, round(montant, 1), classement, statut, source, calcul])

S_RAP25 = "RAP 2025 MIES (extractions/rap2025_execution_par_action_2024_2025.csv)"
S_DPT = "DPT Outre-mer 2026, tableaux par territoire p.248-277 (extractions/passerelle_G5_intrants_RAP_DPT_2024_2025.csv)"
add(2025, "0. Référence DEPP", "État, financement initial, 1er + 2nd degrés (2025p)", etat12_25, "cible", "publié (calcul part x DIE)",
    "NI 26.42 xlsx : Figure 4 B32, C32 ; Figure 5 B32, B33", f"{etat_part_1_25:.6f} x {die1_25:.1f} + {etat_part_2_25:.6f} x {die2_25:.1f}")
add(2025, "A. Budget", "MIES exécutée 2025 (CP, CAS compris)", mies[2025], "", "officiel", "RAP 2025 p.20-21 ; PLRG 2025")
add(2025, "A. Budget", "− Dépenses dans les COM (exécution 2024 reprise pour 2025)", -com_mies_24, "hors champ « France »", "documenté (champ) ; montant 2025 = hypothèse", S_DPT)
add(2025, "A. Budget", "= MIES champ France", france_25, "", "calcul", "")
add(2025, "B. Postes identifiables (RAP)", "− Post-bac en lycée (P141-05 + P139-06)", -postbac[2025], "supérieur (STS, CPGE)", "documenté", S_RAP25 + " ; Dossier 206 p.13-14 imprimées")
add(2025, "B. Postes identifiables (RAP)", "− Formation continue des adultes et VAE (P141-09)", -fc_adultes[2025], "extrascolaire (formation professionnelle continue)", "documenté", S_RAP25 + " ; Dossier 206 p.14")
add(2025, "B. Postes identifiables (RAP)", "− Allocations PFMP (P141 208,0 + P139 54,76 + P143 28,98)", -(pfmp_men[2025] + pfmp_143[2025]), "exclues de la DIE", "documenté", "RAP 2025 p.116, 230, 404 ; L'état de l'École 2025 p.98 imprimée")
add(2025, "B. Postes identifiables (RAP)", "− P214 action 11 (sport, jeunesse, vie associative)", -p214_11[2025], "hors DIE ou extrascolaire", "supposé", S_RAP25)
add(2025, "B. Postes identifiables (RAP)", "− Gratifications des étudiants MEEF (P140, P141, P139, P230)", -meef[2025], "supérieur (master) ou exclues", "supposé", "RAP 2025 p.64, 125, 173, 238-239")
add(2025, "B. Postes identifiables (RAP)", f"− BTSA, BTS et CPGE agricoles du P143 (prorata des effectifs, {100*part_btsa[2025]:.2f} %)", -btsa_25, "supérieur court", "classement documenté ; montant = hypothèse (prorata)",
    "DGER, effectifs voie scolaire 2024 et 2025 (2/3 - 1/3) ; P143 hors COM et hors PFMP", f"{part_btsa[2025]:.5f} x ({p143_france_25:.1f} - {pfmp_143[2025]:.1f})")
add(2025, "B. Postes identifiables (RAP)", "= MIES France après postes identifiables", etape1, "", "calcul", "")
add(2025, "B. Postes identifiables (RAP)", "Écart restant à ce stade (budget − DEPP)", etape1 - etat12_25, "", "calcul", "")
add(2025, "C. Postes estimés (calage DEPP)", f"− Crédits MEN imputés par clés hors 1er/2nd degrés ({100*part_non_ident_24:.3f} % des crédits MEN hors COM)", -c1_central,
    "supérieur (part STS/CPGE des actions transversales et des enseignants imputés aux actions lycée) ; extrascolaire (formation continue des personnels) ; part du P214",
    "principe documenté (Dossier 206 p.14 et 32) ; montant estimé (calage DEPP 2024p)", "voir lignes 2024 ci-dessous", f"{part_non_ident_24:.5f} x {men_france_25:.1f}")
add(2025, "C. Postes estimés (calage DEPP)", "+ Financements de l'État hors MIES en 1er/2nd degrés (autres ministères hors agriculture + reste du monde)", autres_rdm_12_24,
    "État (autres ministères, reste du monde)", "estimé (résidu DEPP 2024p, supposé stable en 2025)", "voir lignes 2024 ci-dessous")
add(2025, "D. Résultat", "= Estimation de la part État DEPP à partir du budget", etape3, "", "calcul", "")
add(2025, "D. Résultat", "Résidu (estimation − DEPP), variante centrale", residu_central, "", "calcul", "|résidu| < 500 MEUR : oui" if abs(residu_central) < 500 else "|résidu| >= 500 MEUR")
add(2025, "E. Variantes", "Résidu si calage complet sur la structure DEPP 2013 (Dossier 206)", var13 - etat12_25, "", "sensibilité", "Tableau 2i 2013 + PLR 2013")
add(2025, "E. Variantes", f"Résidu si BTSA au prorata pondéré par secteur ({100*part_btsa_pond_25:.2f} %)", var_btsa_pond, "", "sensibilité", "")
add(2025, "E. Variantes", f"Résidu si BTSA selon la structure DEPP 2013 ({100*part_btsa_depp13:.2f} %)", var_btsa_d13, "", "sensibilité", "")
add(2025, "E. Variantes", "Résidu si COM 2025 = LFI 2025 (1 103,5 MEUR)", var_com, "", "sensibilité", "DPT 2026")
add(2025, "E. Variantes", "Pour mémoire : ancienne passerelle (MIES − COM − post-bac − P214-11 − FC adultes − apprentissage) − DEPP", ancienne - etat12_25, "", "ancienne version (NOTES § 8)", "")
# 2024 : calage
add(2024, "F. Calage 2024", "État, financement initial, 1er + 2nd degrés (2024p)", etat12_24, "", "publié (calcul)", "NI 25.52 xlsx : Figure 4 B32, C32 ; Figure 5 B31, B32")
add(2024, "F. Calage 2024", "dont MEN-MESR, 1er + 2nd degrés (parts « dont MENESR » arrondies à 0,1 point)", men12_24, "", "publié (calcul)", "séries chronologiques DEPP 366603 : onglets « DIE 1er degré » et « DIE 2nd degré », cellules U17", f"{men_part_1_24} x {die1_24:.1f} + {men_part_2_24} x {die2_24:.1f} (incertitude d'arrondi +/- 66 MEUR)")
add(2024, "F. Calage 2024", "Crédits MEN de la MIES hors COM (2024)", men_france_24, "", "officiel", "RAP 2025 (colonne 2024) ; DPT 2026")
add(2024, "F. Calage 2024", "= Crédits MEN classés par la DEPP hors 1er/2nd degrés ou exclus", hors12_men_24, "", "calcul", "", f"{100*hors12_men_24/men_france_24:.3f} % des crédits MEN hors COM")
add(2024, "F. Calage 2024", "dont postes identifiables 2024 (post-bac 1 570,3 ; FC adultes 92,7 ; PFMP MEN 392,9 ; P214-11 181,0 ; MEEF 13,8)", ident_men_24, "", "officiel (RAP)", "RAP 2024 p.124, 134, 182, 235, 243 ; RAP 2025")
add(2024, "F. Calage 2024", "dont imputations par clés non identifiables", non_ident_24, "", "calcul", "", f"{100*part_non_ident_24:.3f} % des crédits MEN hors COM")
add(2024, "F. Calage 2024", "P143 hors COM, hors PFMP et hors BTSA (part 2nd degré)", p143_2nd_24, "", "estimé (prorata effectifs)", "", f"part supérieur court {100*part_btsa[2024]:.3f} %")
add(2024, "F. Calage 2024", "= Autres ministères (hors agriculture) + reste du monde, 1er + 2nd degrés", autres_rdm_12_24, "", "résidu", "")
add(2013, "G. Calage 2013", "DEPP « Enseignement scolaire » (programmes MEN de la MIES), total", es13["total_MEUR"], "", "publié", "Dossier 206, Tableau 2i 2013 définitif (p.152 imprimée)")
add(2013, "G. Calage 2013", "dont supérieur + extrascolaire", hors12_13, "", "publié", "idem", f"{100*part_hors12_13:.3f} %")
add(2013, "G. Calage 2013", "Budget 2013 : post-bac (P141-05 + P139-06) + FC adultes (P141-09)", postbac13 + fc13, "", "officiel", "PLR 2013 (data.economie.gouv.fr)")
add(2013, "G. Calage 2013", "= imputations par clés non identifiables", non_ident_13, "", "calcul", "", f"{100*part_non_ident_13:.3f} % du total DEPP")
add(2013, "G. Calage 2013", "Budget 2013 : programmes MEN de la MIES (France + COM)", men13_budget, "", "officiel", "PLR 2013", f"écart avec DEPP = {men13_budget - es13['total_MEUR']:.1f} (COM + exclusions)")
add(2013, "G. Calage 2013", "Autres ministères + reste du monde, 1er + 2nd degrés", autres_rdm_12_13, "", "publié", "Tableau 2i 2013")
for y in (2024, 2025):
    add(y, "H. Contrôle global MEN-MESR (tous niveaux)", "Budget MIES-MEN + P150 + P231, hors COM (COM 2024), hors PFMP", glob[y][1], "", "calcul", "RAP ; RERS 2026 10.03 ; DPT 2026 ; RAP PFMP")
    add(y, "H. Contrôle global MEN-MESR (tous niveaux)", "DEPP : MEN-MESR, financement initial, tous niveaux", depp_glob[y], "", "publié", "NI 26.42 xlsx Figure 3bis " + ("E5 (2024 définitif)" if y == 2024 else "F5 (2025p)"))
    add(y, "H. Contrôle global MEN-MESR (tous niveaux)", "Écart DEPP − budget", depp_glob[y] - glob[y][1], "", "calcul", "", f"{100*(depp_glob[y]-glob[y][1])/glob[y][1]:.2f} %")
# informations complementaires (formation des enseignants, stagiaires)
add(2025, "I. Pour information", "Actions de formation des personnels (P140-04 + P141-10 + P139-10)", formation_actions[2025], "rémunération des personnels en formation = dépense d'enseignement ; formation continue -> extrascolaire", "documenté (principe)", S_RAP25, f"dont titre 2 {formation_T2_25:.1f} ; hors titre 2 {formation_HT2_25:.1f}")
add(2025, "I. Pour information", "Hors titre 2 des actions de formation, hors gratifications MEEF (borne basse des coûts d'organisation de la formation)", formation_HT2_25 - (I['meef_p140_2025'] + I['meef_p141_2025'] + I['meef_p139_2025']) / M, "", "calcul", S_RAP25)
for p in (140, 141, 230, 139):
    e, c, h, w = stag[p]
    add(2025, "I. Pour information", f"Stagiaires P{p} (catégorie 1108) : {e:.2f} ETPT x {int(c)} € (hors CAS)", h, "dépense d'enseignement ; niveau non documenté", "officiel (ETPT, coût) ; produit = calcul", "RAP 2025", f"avec CAS (estimation) : {w:.1f}")
add(2025, "I. Pour information", "Stagiaires, total 4 programmes (hors CAS)", sum(v[2] for v in stag.values()), "", "calcul", "RAP 2025", f"avec CAS (estimation) : {sum(v[3] for v in stag.values()):.1f} ; {sum(v[0] for v in stag.values()):.2f} ETPT")

out = os.path.join(HERE, "_derive_G5_passerelle_MIES_DEPP_2024_2025.csv")
with open(out, "w", newline="", encoding="utf-8-sig") as f:
    w = csv.writer(f, delimiter=";")
    w.writerow(["annee", "bloc", "poste", "montant_MEUR", "classement_DEPP", "statut", "source_emplacement", "calcul_remarque"])
    w.writerows(rows)

# ---------------------------------------------------------------- resume console
print(f"DEPP Etat 1+2 2025p = {etat12_25:.1f} ; 2024p = {etat12_24:.1f} ; MEN-MESR 1+2 2024p = {men12_24:.1f}")
print(f"MIES 2025 = {mies[2025]:.1f} ; MEN = {men[2025]:.1f} ; P143 = {p143[2025]:.1f} ; COM 2024 = {com_mies_24:.1f} (MEN {com_men_24:.1f}) ; COM MIRES 2024 = {com_mires_24:.1f}")
print(f"Part BTSA (annee civile) 2024 = {100*part_btsa[2024]:.3f} % ; 2025 = {100*part_btsa[2025]:.3f} % ; ponderee 2025 = {100*part_btsa_pond_25:.2f} % ; DEPP 2013 = {100*part_btsa_depp13:.2f} %")
print(f"Postes identifiables 2025 = {ident_25:.1f} -> etape 1 = {etape1:.1f} (ecart {etape1-etat12_25:.1f})")
print(f"Calage : non identifie 2024 = {non_ident_24:.1f} ({100*part_non_ident_24:.3f} %) ; 2013 = {non_ident_13:.1f} ({100*part_non_ident_13:.3f} %)")
print(f"C1 2025 central = {c1_central:.1f} (variante 2013 : {c1_var13:.1f}) ; autres ministeres + RDM 2024 = {autres_rdm_12_24:.1f} (2013 : {autres_rdm_12_13:.1f})")
print(f"Residu central 2025 = {residu_central:.1f} ; var. 2013 = {var13-etat12_25:.1f} ; BTSA pond. = {var_btsa_pond:.1f} ; BTSA DEPP13 = {var_btsa_d13:.1f} ; COM LFI = {var_com:.1f} ; ancienne passerelle = {ancienne-etat12_25:.1f}")
for y in (2024, 2025):
    print(f"Controle global {y} : budget hors PFMP {glob[y][1]:.1f} (avec PFMP {glob[y][0]:.1f}) ; DEPP {depp_glob[y]:.1f} ; ecart {depp_glob[y]-glob[y][1]:.1f}")
print(f"2013 : DEPP ES {es13['total_MEUR']:.1f} ; budget MEN {men13_budget:.1f} ; hors 1+2 {hors12_13:.1f} ({100*part_hors12_13:.3f} %) ; post-bac {postbac13:.1f} ; FC {fc13:.1f} ; part 2nd degre P143 {100*agri13_2nd_share:.2f} %")
print(f"Stagiaires : {sum(v[0] for v in stag.values()):.2f} ETPT ; HCAS {sum(v[2] for v in stag.values()):.1f} ; avec CAS (est.) {sum(v[3] for v in stag.values()):.1f}")
print(f"Formation actions 2025 = {formation_actions[2025]:.1f} (T2 {formation_T2_25:.1f}, HT2 {formation_HT2_25:.1f}) ; 2024 = {formation_actions[2024]:.1f}")
print("ecrit", out)
